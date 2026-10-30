import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from materials_word import make_materials_doc


class MaterialsWordTest(unittest.TestCase):
    def test_checklist_is_editable_and_preserves_all_rows_and_warnings(self):
        items = [dict(category='附件待关联', item=f'第{i}份路径图', state='未关联文件，不代表资料未提供', action='核对已有文件位置和路径版本', blocking=i == 0) for i in range(35)]
        data = dict(title='协议征询函资料核对与补充清单', projectName='中文测试工程', routeId='P01', version='v1', stale=True, simulation=False, exportedAt='2026-10-30T00:00:00Z', itemCount=len(items), groups=[dict(agreementIds=['L-01'], unit='测试接收单位', status='failed', error='资料需核对', items=items)])
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / 'checklist.docx'
            make_materials_doc(data, output)
            d = Document(output)
            self.assertEqual(d.paragraphs[0].style.name, 'Title')
            self.assertFalse(list(d.styles['Title'].element.iter(qn('w:pBdr'))))
            text = '\n'.join(p.text for p in d.paragraphs)
            self.assertIn('上一次已同步', text)
            self.assertIn('生成或检查失败', text)
            self.assertIn('失败原因：资料需核对', text)
            self.assertIn('2026-10-30 08:00', text)
            self.assertEqual(len(d.tables[0].rows), 36)
            self.assertIn('第34份路径图', d.tables[0].rows[-1].cells[1].text)
            self.assertIn('阻断生成', d.tables[0].rows[1].cells[2].text)
            self.assertTrue(list(d.tables[0].rows[0]._tr.iter(qn('w:tblHeader'))))
            self.assertLess(abs(sum(column.width for column in d.tables[0].columns) - 6120000), 635)

    def test_no_pending_fields_does_not_claim_formal_approval(self):
        data = dict(title='协议征询函资料核对与补充清单', projectName='测试工程', routeId='P01', version='v1', itemCount=0, groups=[dict(agreementIds=['L-02'], unit='测试单位', status='generating', items=[])])
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / 'empty.docx'
            make_materials_doc(data, output)
            text = '\n'.join(p.text for p in Document(output).paragraphs)
            self.assertIn('正文生成中', text)
            self.assertIn('仍需按流程核对正文', text)
            self.assertNotIn('检查通过', text)


if __name__ == '__main__':
    unittest.main()
