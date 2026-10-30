"""Reference-derived letter layout shared by single DOCX and ZIP exports."""
from copy import deepcopy
from pathlib import Path
import re
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.enum.style import WD_STYLE_TYPE

TEMPLATES = Path(__file__).with_name('templates')
BODY_FONT = '仿宋_GB2312'
TITLE_FONT = '方正小标宋_GBK'

def balanced_title(text, capacity):
    """Balance long Chinese titles without splitting Latin codes or dropping text."""
    if '\n' in text:
        return text
    weight = lambda char: .5 if ord(char) < 128 else 1
    total = sum(map(weight, text))
    lines = max(1, int((total + capacity - .001) // capacity))
    result = []
    for remaining in range(lines, 1, -1):
        target = sum(map(weight, text)) / remaining
        candidates = []
        width = 0
        for index in range(1, len(text)):
            width += weight(text[index-1])
            if text[index-1].isascii() and text[index-1].isalnum() and text[index].isascii() and text[index].isalnum():
                continue
            penalty = 8 if text[index:].startswith('的函') else 0
            bonus = 1 if text[index-1].isspace() or text[index].isspace() or text[:index].endswith(('工程','线路','系统','方案')) else 0
            candidates.append((abs(width-target)*2+penalty-bonus,index))
        if not candidates:
            break
        index = min(candidates)[1]
        result.append(text[:index].rstrip());text=text[index:].lstrip()
    return '\n'.join([*result,text])


def font(run, name=BODY_FONT, size=16, bold=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    props = run._element.get_or_add_rPr()
    props.rFonts.set(qn('w:eastAsia'), name)
    props.rFonts.set(qn('w:cs'), 'Times New Roman')
    for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        props.rFonts.attrib.pop(qn('w:' + attr), None)
    for tag in ['highlight', 'shd', 'ins', 'del']:
        for child in list(props.findall(qn('w:' + tag))):
            props.remove(child)


def make_doc(letter, target):
    # All outgoing consultation letters follow the user-selected plain DOC.
    # Red-headed reply PDFs are incoming originals, not outgoing DOCX templates.
    d = Document(TEMPLATES / 'letter-plain.docx')
    patterns = {p.text[2:-2]: deepcopy(p._p) for p in d.paragraphs if p.text.startswith('{{')}
    for child in list(d._element.body):
        if child.tag != qn('w:sectPr'):
            d._element.body.remove(child)
    d.core_properties.title = letter['title']
    d.core_properties.author = ''
    d.core_properties.last_modified_by = ''
    normal = d.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(16)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), BODY_FONT)
    if 'Title' not in d.styles:
        d.styles.add_style('Title', WD_STYLE_TYPE.PARAGRAPH)

    def paragraph(text, role='body', *, size=16, bold=False):
        element = deepcopy(patterns[role])
        for child in list(element):
            if child.tag != qn('w:pPr'):
                element.remove(child)
        d._element.body.insert(len(d._element.body) - 1, element)
        p = Paragraph(element, d._body)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.widow_control = True
        # Reuse source roles, but normalize legacy character indents and tabs.
        props = p._p.get_or_add_pPr()
        for tag in ['ind', 'tabs', 'numPr', 'outlineLvl']:
            node = props.find(qn('w:' + tag))
            if node is not None:
                props.remove(node)
        p.paragraph_format.line_spacing = Pt(28)
        p.alignment = 3 if role == 'body' else 0
        if role == 'body':
            p.paragraph_format.first_line_indent = Pt(32)
        run = p.add_run(str(text))
        font(run, size=size, bold=bold)
        return p

    section = d.sections[0]
    title_capacity = (section.page_width-section.left_margin-section.right_margin)/Pt(18)
    title = paragraph(balanced_title(letter['title'], title_capacity), 'title')
    title.style = d.styles['Title']
    font(title.runs[0], TITLE_FONT, 18, bold=True)
    title.alignment = 1
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.left_indent = Pt(0)
    title.paragraph_format.right_indent = Pt(0)
    # Reference-sized fixed leading avoids different fallback-font ascenders
    # doubling the title leading in Word/LibreOffice on another device.
    title.paragraph_format.line_spacing = Pt(28)
    title.paragraph_format.space_after = Pt(14)
    title.paragraph_format.keep_with_next = True

    if letter['issuer'].get('documentNumber'):
        p = paragraph(letter['issuer']['documentNumber'], 'recipient')
        p.alignment = 1
        p.paragraph_format.keep_with_next = True
    p = paragraph(letter['unit'] + '：', 'recipient')
    p.paragraph_format.keep_with_next = True
    for text in letter['bodyBlocks']:
        paragraph(text)

    printable_cm = (d.sections[0].page_width - d.sections[0].left_margin - d.sections[0].right_margin) / Cm(1)
    for t in letter['tables']:
        caption = paragraph(t['title'], 'caption')
        caption.paragraph_format.keep_with_next = True
        caption.paragraph_format.space_before = Pt(8)
        caption.paragraph_format.space_after = Pt(4)
        count = len(t['headers'])
        table = d.add_table(rows=1, cols=count)
        table.autofit = False
        table.alignment = 1
        borders = OxmlElement('w:tblBorders')
        for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            edge = OxmlElement('w:' + side)
            for key, value in {'val': 'single', 'sz': '4', 'color': '000000'}.items():
                edge.set(qn('w:' + key), value)
            borders.append(edge)
        table._tbl.tblPr.append(borders)
        margins = OxmlElement('w:tblCellMar')
        for side in ['top', 'left', 'bottom', 'right']:
            edge = OxmlElement('w:' + side)
            edge.set(qn('w:w'), '40')
            edge.set(qn('w:type'), 'dxa')
            margins.append(edge)
        table._tbl.tblPr.append(margins)
        repeat = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for c, text in zip(table.rows[0].cells, t['headers']):
            c.text = str(text)
        for row in t['rows']:
            for c, text in zip(table.add_row().cells, row):
                c.text = str(text)
        weights = [1, 2.3, 2.3, 1.3, 1.5, 2, 2.5, 3] if count == 8 else [1] * count
        widths = [printable_cm * weight / sum(weights) for weight in weights]
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                cell.vertical_alignment = 1
                for p in cell.paragraphs:
                    p.alignment = 1
                    p.paragraph_format.first_line_indent = Pt(0)
                    p.paragraph_format.line_spacing = Pt(15)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    for run in p.runs:
                        font(run, '宋体', 12)
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))

    closing = []
    if letter['attachments']:
        p = paragraph('附件：', 'attachment')
        closing.append(p)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.keep_with_next = True
        for a in letter['attachments']:
            p = paragraph(f"{a['number']}. {a['name']}" + ('' if a['available'] else '（文件待关联/核对）'), 'attachment')
            closing.append(p)
            p.paragraph_format.left_indent = Pt(32)
            p.paragraph_format.first_line_indent = Pt(0)
            # Hanging numbering avoids long attachment names starting at the margin.
            p.paragraph_format.first_line_indent = Pt(-16)

    issuer = letter['issuer']
    p = paragraph(issuer.get('issuer') or '发函单位待填写', 'signature')
    closing.append(p)
    p.alignment = 2
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.keep_with_next = True
    date = issuer.get('issueDate') or '日期待填写'
    match = re.fullmatch(r'(\d{4})-(\d{2})-(\d{2})', date)
    if match:
        date = f'{match[1]}年{int(match[2])}月{int(match[3])}日'
    p = paragraph(date, 'date')
    closing.append(p)
    p.alignment = 2
    # Keep the small signature/contact block intact across page boundaries.
    p.paragraph_format.keep_with_next = True
    p = paragraph('联系人：' + (issuer.get('contactName') or '待填写') + '　联系电话：' + (issuer.get('contactPhone') or '待填写'), 'contact')
    closing.append(p)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.keep_with_next = True
    closing.append(paragraph('复函地址：' + (issuer.get('replyAddress') or '待填写'), 'replyAddress'))
    # A short closing block travels together, so a page does not contain only
    # an isolated signature/contact. Large attachment lists remain free to flow.
    if len(closing) <= 9 and sum(len(p.text) for p in closing) <= 300:
        for p in closing[:-1]:
            p.paragraph_format.keep_with_next = True

    # The exported letter contains only letter content. Draft/simulation labels
    # belong to filenames; validation and review notes stay in their own records.
    d.save(target)
    # Keep source footer drawings and PAGE fields byte-for-byte. python-docx
    # otherwise rewrites their namespaces even though they were never edited.
    with ZipFile(TEMPLATES / 'letter-plain.docx') as template:
        chrome = {name: template.read(name) for name in template.namelist()
                  if name.startswith(('word/header', 'word/footer'))}
    with ZipFile(target) as generated:
        entries = [(info, chrome.get(info.filename, generated.read(info.filename)))
                   for info in generated.infolist()]
    with ZipFile(target, 'w', ZIP_DEFLATED) as output:
        for info, content in entries:
            output.writestr(info, content)
