"""Extract reference text only. Original files and layouts are retained separately."""
import json,sys,subprocess
from pathlib import Path

p=Path(sys.argv[1]);suffix=p.suffix.lower()
if suffix=='.pdf':
    from pypdf import PdfReader
    reader=PdfReader(p)
    if len(reader.pages)>200:raise ValueError('PDF超过200页，请按章节拆分')
    text='\n\n'.join('【第%d页】\n%s'%(i+1,page.extract_text() or '') for i,page in enumerate(reader.pages))
elif suffix=='.docx':
    from docx import Document
    d=Document(p)
    text='\n'.join([x.text for x in d.paragraphs]+[' | '.join(c.text for c in row.cells) for table in d.tables for row in table.rows])
elif suffix=='.doc' and sys.platform=='darwin':
    text=subprocess.check_output(['/usr/bin/textutil','-convert','txt','-stdout',str(p)],timeout=30).decode('utf-8')
elif suffix in ('.txt','.md'):
    raw=p.read_bytes()
    try:text=raw.decode('utf-8-sig')
    except UnicodeDecodeError:text=raw.decode('gb18030')
else:raise ValueError('支持PDF、Word、TXT、Markdown；旧DOC文件仅Mac支持解析')
if len(text)>80000:raise ValueError('文字超过8万字，请拆分章节')
print(json.dumps({'text':text},ensure_ascii=False))
