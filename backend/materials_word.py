"""Create a genuine DOCX checklist from the task's structured material checks."""
from datetime import datetime
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def font(run, size=11, bold=False, family='宋体'):
    run.font.name = family
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), family)


def paragraph(document, text, size=11, bold=False, style=None):
    p = document.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = Pt(max(16, size * 1.3))
    font(p.add_run(str(text)), size, bold, '黑体' if bold else '宋体')
    return p


def make_materials_doc(data, output):
    d = Document()
    section = d.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    d.core_properties.title = data['title']
    d.core_properties.author = ''
    normal = d.styles['Normal']
    normal.font.name, normal.font.size = '宋体', Pt(11)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    for style in d.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for name in ['Title', 'Heading 1']:
        d.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title = paragraph(d, data['title'], 20, True, 'Title')
    title.paragraph_format.space_after = Pt(14)
    for label, value in [('工程', data['projectName']), ('路径', data['routeId']), ('数据版本', data['version'])]:
        paragraph(d, label + '：' + (value or '待填写'))
    if data.get('exportedAt'):
        stamp = datetime.fromisoformat(data['exportedAt'].replace('Z', '+00:00')).astimezone(ZoneInfo('Asia/Shanghai'))
        paragraph(d, '导出时间：' + stamp.strftime('%Y-%m-%d %H:%M') + '（北京时间）')
    paragraph(d, f"本清单共涉及 {len(data['groups'])} 份函件，列出 {data['itemCount']} 项待处理资料。请先核对现有文件和平台记录，确实缺少时再补充。未填写或未关联不等于甲方未提供。")
    if data.get('stale'):
        paragraph(d, '注意：当前输入有资料待同步，本清单依据上一次已同步的任务生成。', bold=True)
    if data.get('simulation'):
        paragraph(d, '模拟测试清单：模拟资料及附件不作为真实工程资料使用。', bold=True)
    status_names = {'pending': '等待生成', 'generating': '正文生成中', 'checking': '正文检查中', 'failed': '生成或检查失败', 'blocked': '存在阻断项', 'incomplete': '资料待处理', 'needs_rewrite': '正文待更新', 'draft': '草稿待复核', 'confirmed': '已人工复核', 'simulated': '已模拟确认'}
    widths = [1.8, 3, 4, 8.2]
    for index, group in enumerate(data['groups'], 1):
        heading = paragraph(d, f"{index} {group['unit']}", 14, True, 'Heading 1')
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.keep_with_next = True
        p = paragraph(d, '协议编号：' + ('、'.join(group['agreementIds']) or '待填写') + '    函件状态：' + status_names.get(group['status'], group['status'] or '待核对'), 10.5)
        p.paragraph_format.keep_with_next = True
        if group.get('error'):
            p = paragraph(d, '失败原因：' + group['error'])
            p.paragraph_format.keep_with_next = bool(group['items'])
        if not group['items']:
            paragraph(d, '当前任务记录没有列出待处理资料；仍需按流程核对正文及发出状态。')
            continue
        table = d.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        for cell, label in zip(table.rows[0].cells, ['类别', '具体项目', '当前情况', '处理方式']):
            cell.text = label
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'F2F2F2')
            cell._tc.get_or_add_tcPr().append(shade)
        repeat = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for item in group['items']:
            row = table.add_row()
            state = item['state'] + ('；当前阻断生成' if item['blocking'] else '')
            for cell, text in zip(row.cells, [item['category'], item['item'], state, item['action']]):
                cell.text = text
        for row_index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.line_spacing = Pt(15)
                    p.paragraph_format.keep_with_next = row_index == 0
                    for run in p.runs:
                        font(run, 10.5, row_index == 0)
    paragraph(d, '核对反馈可注明：已提供文件的名称和位置、适用路径版本；确需补交的资料内容及联系人。')
    d.save(output)
