"""Contract checks for the public DOCX/ZIP renderer (no model/network calls)."""
import unittest, tempfile, json, subprocess, sys, io
from pathlib import Path
from zipfile import ZipFile
from copy import deepcopy
from lxml import etree as E
from docx import Document
from word_layout import make_doc

NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

def fixture(kind='general'):
    return dict(id='FORMAT-TEST',agreementIds=['TEST-1'],round='first',template='版式测试 v1.0',unit='测试收函单位',type=kind,typeLabel='测试类型',title='关于征询测试工程路径方案意见的函',
        issuer=dict(name='测试工程',issuer='本次测试发函单位',issueDate='2026-10-30',contactName='测试联系人',contactPhone='测试电话',replyAddress='测试地址'),
        bodyBlocks=['测试工程概况，正文应完整保留。','本次征询事项应完整保留。'],tables=[],
        attachments=[dict(number=1,name='测试路径图.pdf',available=False)],status='draft',missingFields=[dict(label='路径图实际文件')],warnings=['待业务复核','当前工程为演示数据，正文及参数仅供演示。'],scope='测试区段',revision=1,confirmations=[])


def xml(path):
    with ZipFile(path) as z:return E.fromstring(z.read('word/document.xml'))


class WordLayoutTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()
    def render(self,letter,name='letter.docx'):
        target=self.root/name;make_doc(letter,target);return target,xml(target)
    def test_consultation_uses_plain_title_and_current_signature(self):
        path,tree=self.render(fixture())
        self.assertFalse(tree.xpath('.//w:rPr/w:color[@w:val="FF0000"]|.//w:pBdr|.//w:drawing|.//w:pict',namespaces=NS))
        body=tree.xpath('.//w:p[w:r/w:t="测试工程概况，正文应完整保留。"]',namespaces=NS)[0]
        self.assertEqual(body.find('w:r/w:rPr/w:sz',NS).get('{'+NS['w']+'}val'),'32')
        self.assertEqual(body.find('w:r/w:rPr/w:rFonts',NS).get('{'+NS['w']+'}eastAsia'),'仿宋_GB2312')
        title=Document(path).paragraphs[0]
        self.assertEqual(title.style.name,'Title');self.assertEqual(title.runs[0].font.size.pt,18)
        self.assertEqual(title.runs[0]._element.rPr.rFonts.get('{'+NS['w']+'}eastAsia'),'方正小标宋_GBK')
        text=''.join(tree.xpath('//w:t/text()',namespaces=NS))
        self.assertEqual(text.count('本次测试发函单位'),1)
        self.assertIn('2026年10月30日',text)
        for notice in ['草稿 · 待人工复核','待补资料与复核事项','当前工程为演示数据','路径图实际文件','待业务复核']:
            self.assertNotIn(notice,text)
        with ZipFile(path) as z:
            allxml=''.join(z.read(n).decode() for n in z.namelist() if n.endswith('.xml'))
            for old in ['广州局','历史联系人甲','历史联系电话甲','历史联系人乙','历史联系电话乙','历史联系人丙','中国能源建设集团广东省电力设计研究院有限公司']:
                self.assertNotIn(old,allxml)
            footers=''.join(z.read(n).decode() for n in z.namelist() if n.startswith('word/footer') and n.endswith('.xml'))
            self.assertNotIn('FF0000',footers)
        paragraphs=Document(path).paragraphs
        self.assertEqual(paragraphs[-1].text,'复函地址：测试地址')
        self.assertFalse(any('待人工复核' in p.text for p in paragraphs))
        signature=next(p for p in paragraphs if p.text=='本次测试发函单位')
        date=next(p for p in paragraphs if p.text=='2026年10月30日')
        for p in [signature,date]:
            self.assertEqual(p.alignment,2)
            self.assertFalse(p._p.xpath('./w:pPr/w:ind[@w:endChars or @w:startChars]'))
        self.assertFalse(tree.xpath('.//w:ins|.//w:del|.//w:highlight',namespaces=NS))
    def test_all_consultation_types_have_no_red_letterhead(self):
        for kind in ['general','grounding','mineral','water','highway','aviation','railway','seismic','military']:
            with self.subTest(kind=kind):
                path,tree=self.render(fixture(kind))
                self.assertFalse(tree.xpath('.//w:rPr/w:color[@w:val="FF0000"]|.//w:pBdr|.//w:pict|.//w:drawing',namespaces=NS))
                self.assertEqual(Document(path).paragraphs[0].text,fixture(kind)['title'])
    def test_plain_specialty_table_all_rows_and_printable_width(self):
        letter=fixture('highway');letter['tables']=[dict(title='两档跨越参数表',headers=['序号','设施','跨越档','交角','档距','水平距离','净空','里程'],rows=[['1','测试高速','NA1～NA2','78°','168','47.34','29','高速匝道'],['2','测试高速','NA2～NA3','80°','688','78.05','27.5','K24+750m']])]
        path,tree=self.render(letter);d=Document(path)
        self.assertFalse(tree.xpath('.//w:color[@w:val="FF0000"]|.//w:pBdr',namespaces=NS))
        self.assertEqual(d.paragraphs[0].runs[0].font.size.pt,18)
        self.assertEqual(len(d.tables[0].rows),3)
        self.assertEqual(d.tables[0].cell(2,6).text,'27.5')
        self.assertTrue(tree.find('.//w:tblHeader',NS) is not None)
        self.assertEqual(len(tree.findall('.//w:cantSplit',NS)),3)
        width=sum(int(n.get('{'+NS['w']+'}w')) for n in tree.findall('.//w:tblGrid/w:gridCol',NS))
        sec=d.sections[0];expected=(sec.page_width-sec.left_margin-sec.right_margin)/635
        self.assertLessEqual(abs(width-expected),5)
    def test_simulated_and_confirmed_states(self):
        l=fixture();l.update(simulation=True,status='simulated')
        _,tree=self.render(l)
        self.assertNotIn('非真实工程资料',''.join(tree.xpath('//w:t/text()',namespaces=NS)))
        l.pop('simulation');l.update(status='confirmed',missingFields=[],warnings=[])
        _,tree=self.render(l)
        text=''.join(tree.xpath('//w:t/text()',namespaces=NS))
        self.assertNotIn('草稿',text);self.assertNotIn('待补资料与复核事项',text)
    def test_zip_and_single_document_same_content(self):
        letters=[fixture(),fixture('highway')];out=self.root/'bundle.zip'
        original=b'%PDF-1.3\noriginal-incoming-reply-bytes\n%%EOF'
        reply=self.root/'reply.pdf';reply.write_bytes(original)
        letters[0]['attachments'].append(dict(number=2,name='复函原件.pdf',available=True,assetId='reply'))
        result=subprocess.run([sys.executable,str(Path(__file__).with_name('export_letters.py'))],input=json.dumps(dict(output=str(out),format='zip',letters=letters,assets={'reply':{'path':str(reply),'name':'复函原件.pdf'}})),capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)
        with ZipFile(out) as z:
            docs=[n for n in z.namelist() if n.endswith('.docx')]
            self.assertEqual(len(docs),2);self.assertIn('协议征询台账.xlsx',z.namelist())
            self.assertTrue(all(Path(n).name.startswith('草稿_') for n in docs))
            self.assertEqual(z.read(next(n for n in z.namelist() if n.endswith('复函原件.pdf'))),original)
            for name,letter in zip(docs,letters):
                path,_=self.render(letter)
                with ZipFile(io.BytesIO(z.read(name))) as a,ZipFile(path) as b:
                    self.assertEqual(a.read('word/document.xml'),b.read('word/document.xml'))
                    text=''.join(E.fromstring(a.read('word/document.xml')).xpath('//w:t/text()',namespaces=NS))
                    for notice in ['待补资料与复核事项','当前工程为演示数据','草稿 · 待人工复核']:
                        self.assertNotIn(notice,text)

if __name__=='__main__':unittest.main()
